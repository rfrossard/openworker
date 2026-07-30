from __future__ import annotations

import zipfile

from PIL import Image
from pypdf import PdfReader
from pptx import Presentation

from coworker.server.manager import SessionManager
from coworker.skills import SkillLoader
from coworker.tools.presentation_generation import make_build_presentation_tool


def _sample_image(path):
    Image.new("RGB", (640, 360), (35, 76, 142)).save(path, "PNG")


def test_build_presentation_creates_real_matching_files_with_image(tmp_path):
    image = tmp_path / "visual.png"
    _sample_image(image)
    tool = make_build_presentation_tool(workspace=tmp_path)

    result = tool(
        title="Evidence into decisions",
        subtitle="A validated presentation pipeline",
        slides=[
            {
                "title": "Visual evidence improves the narrative",
                "takeaway": "One structured specification drives both final formats.",
                "bullets": ["Research claims first", "Use a distinct relevant visual"],
                "image_path": "visual.png",
                "image_caption": "Test visual",
                "layout": "image-left",
                "image_fit": "contain",
                "image_focus": "right",
                "image_required": True,
                "sources": ["https://example.com/research"],
            },
            {
                "title": "Quality gates prevent fake deliverables",
                "layout": "statement",
                "bullets": ["Validate signatures", "Keep output editable"],
            },
        ],
        pptx_path="reports/deck.pptx",
        pdf_path="reports/deck.pdf",
        minimum_images=1,
    )

    assert result["ok"] is True
    assert result["slides"] == 3
    assert result["images_embedded"] == 1
    assert result["visual_plan_complete"] is True
    assert len(result["preview_paths"]) == 3
    assert result["visual_review_required"] is True
    pptx = tmp_path / "reports/deck.pptx"
    pdf = tmp_path / "reports/deck.pdf"
    assert pptx.read_bytes().startswith(b"PK")
    assert pdf.read_bytes().startswith(b"%PDF-")
    assert len(Presentation(pptx).slides) == 3
    assert len(PdfReader(pdf).pages) == 3
    for path in result["preview_paths"]:
        assert (tmp_path / path).read_bytes().startswith(b"\x89PNG")
    assert (tmp_path / result["contact_sheet_path"]).read_bytes().startswith(b"\x89PNG")
    with zipfile.ZipFile(pptx) as archive:
        assert any(name.startswith("ppt/media/") for name in archive.namelist())


def test_build_presentation_supports_designer_layouts(tmp_path):
    tool = make_build_presentation_tool(workspace=tmp_path)
    result = tool(
        title="Designed from Markdown",
        slides=[
            {
                "title": "A new chapter",
                "takeaway": "The section layout creates a deliberate narrative break.",
                "layout": "section",
            },
            {
                "title": "A voice worth featuring",
                "takeaway": "The quote layout makes one message visually dominant.",
                "layout": "quote",
                "bullets": ["Source attribution"],
            },
            {
                "title": "Two sides of the decision",
                "takeaway": "The content remains editable.",
                "layout": "two-column",
                "bullets": ["Benefit one", "Benefit two", "Risk one", "Risk two"],
            },
        ],
        pptx_path="designer.pptx",
        pdf_path="designer.pdf",
    )

    assert result["ok"] is True
    assert len(Presentation(tmp_path / "designer.pptx").slides) == 4
    assert len(PdfReader(tmp_path / "designer.pdf").pages) == 4


def test_build_presentation_supports_extended_designer_layouts(tmp_path):
    image = tmp_path / "visual.png"
    _sample_image(image)
    layouts = [
        "title-only", "big-number", "checklist", "timeline", "process", "comparison",
        "pros-cons", "three-columns", "four-cards", "metric-grid", "image-background",
        "image-top", "image-bottom", "agenda", "conclusion",
    ]
    image_layouts = {"image-background", "image-top", "image-bottom"}
    tool = make_build_presentation_tool(workspace=tmp_path)
    result = tool(
        title="Extended designer layouts",
        slides=[
            {
                "title": f"{layout} composition",
                "takeaway": "A deliberate visual hierarchy",
                "bullets": ["42%", "First idea", "Second idea", "Third idea", "Fourth idea"],
                "layout": layout,
                "image_path": "visual.png" if layout in image_layouts else "",
                "image_required": layout in image_layouts,
            }
            for layout in layouts
        ],
        pptx_path="extended-designer.pptx",
        pdf_path="extended-designer.pdf",
        minimum_images=3,
        template_id="midnight",
    )

    assert result["ok"] is True
    assert result["images_embedded"] == 3
    assert len(Presentation(tmp_path / "extended-designer.pptx").slides) == 16
    assert len(PdfReader(tmp_path / "extended-designer.pdf").pages) == 16


def test_build_presentation_supports_editable_semantic_visuals(tmp_path):
    tool = make_build_presentation_tool(workspace=tmp_path)
    result = tool(
        title="Semantic visual language",
        slides=[
            {"title": "A comparison table", "layout": "table", "bullets": ["Option | Cost | Speed", "A | 10 | Fast", "B | 15 | Medium"], "sources": ["https://example.com/table"]},
            {"title": "Growth by segment", "layout": "bar-chart", "bullets": ["Core | 72", "New | 44", "Partner | 28"], "sources": ["https://example.com/chart"]},
            {"title": "Revenue mix", "layout": "donut-chart", "bullets": ["Product | 55", "Services | 30", "Other | 15"], "sources": ["https://example.com/mix"]},
            {"title": "Decision flow", "layout": "flow-diagram", "bullets": ["Discover", "Validate", "Build", "Measure"]},
            {"title": "Accountable team", "layout": "org-chart", "bullets": ["CEO > Product", "CEO > Engineering", "CEO > Sales"]},
            {"title": "Delivery roadmap", "layout": "roadmap", "bullets": ["Q1 Research", "Q2 Pilot", "Q3 Launch", "Q4 Scale"]},
        ],
        pptx_path="semantic.pptx",
        pdf_path="semantic.pdf",
    )

    assert result["ok"] is True
    assert result["quality_gate"]["passed"] is True
    assert result["quality_gate"]["score"] == 100
    presentation = Presentation(tmp_path / "semantic.pptx")
    assert len(presentation.slides) == 7
    assert any(shape.has_table for shape in presentation.slides[1].shapes)
    assert presentation.slides[2].has_notes_slide
    assert "[Sources]" in presentation.slides[2].notes_slide.notes_text_frame.text
    assert any(getattr(shape, "has_chart", False) for shape in presentation.slides[2].shapes)
    assert any(getattr(shape, "has_chart", False) for shape in presentation.slides[3].shapes)
    bar_chart = next(shape.chart for shape in presentation.slides[2].shapes if getattr(shape, "has_chart", False))
    assert bar_chart.chart_type == 57  # BAR_CLUSTERED
    assert bar_chart.plots[0].has_data_labels is True
    assert len(PdfReader(tmp_path / "semantic.pdf").pages) == 7


def test_build_presentation_rejects_incomplete_semantic_data(tmp_path):
    tool = make_build_presentation_tool(workspace=tmp_path)
    chart = tool(
        title="Invalid chart",
        slides=[{"title": "Not enough data", "layout": "bar-chart", "bullets": ["Only | 1"]}],
        pptx_path="invalid-chart.pptx",
        pdf_path="invalid-chart.pdf",
    )
    table = tool(
        title="Invalid table",
        slides=[{"title": "Not enough rows", "layout": "table", "bullets": ["A | B"]}],
        pptx_path="invalid-table.pptx",
        pdf_path="invalid-table.pdf",
    )
    assert chart["ok"] is False
    assert "at least two chart rows" in chart["error"]
    assert table["ok"] is False
    assert "header and at least one" in table["error"]


def test_build_presentation_rejects_path_escape_and_missing_image(tmp_path):
    tool = make_build_presentation_tool(workspace=tmp_path)
    escaped = tool(
        title="Deck",
        slides=[{"title": "Slide"}],
        pptx_path="../deck.pptx",
        pdf_path="deck.pdf",
    )
    missing = tool(
        title="Deck",
        slides=[{"title": "Slide", "image_path": "missing.png"}],
        pptx_path="deck.pptx",
        pdf_path="deck.pdf",
    )
    assert escaped["ok"] is False
    assert "escapes" in escaped["error"]
    assert missing["ok"] is False
    assert "not found" in missing["error"]

    incomplete = tool(
        title="Deck",
        slides=[{"title": "Slide without its planned asset"}],
        pptx_path="incomplete.pptx",
        pdf_path="incomplete.pdf",
        minimum_images=1,
    )
    required = tool(
        title="Deck",
        slides=[{"title": "Required visual", "image_required": True}],
        pptx_path="required.pptx",
        pdf_path="required.pdf",
    )
    assert incomplete["ok"] is False
    assert "requires at least 1 images" in incomplete["error"]
    assert required["ok"] is False
    assert "requires its planned image" in required["error"]


def test_build_presentation_applies_builtin_and_workspace_potx_templates(tmp_path):
    tool = make_build_presentation_tool(workspace=tmp_path)
    built_in = tool(
        title="Template catalog",
        slides=[{"title": "Editable theme"}],
        pptx_path="built-in.pptx",
        pdf_path="built-in.pdf",
        template_id="forest",
    )
    assert built_in["ok"] is True
    assert built_in["template_id"] == "forest"
    assert built_in["template_path"] is None

    source_pptx = tmp_path / "brand-source.pptx"
    potx = tmp_path / "brand.potx"
    Presentation().save(source_pptx)
    with zipfile.ZipFile(source_pptx) as source, zipfile.ZipFile(
        potx, "w", compression=zipfile.ZIP_DEFLATED
    ) as output:
        for name in source.namelist():
            data = source.read(name)
            if name == "[Content_Types].xml":
                data = data.replace(
                    b"application/vnd.openxmlformats-officedocument.presentationml.presentation.main+xml",
                    b"application/vnd.openxmlformats-officedocument.presentationml.template.main+xml",
                )
            output.writestr(name, data)
    custom = tool(
        title="Brand template",
        slides=[{"title": "Inherited master"}],
        pptx_path="custom.pptx",
        pdf_path="custom.pdf",
        template_path="brand.potx",
    )
    assert custom["ok"] is True
    assert custom["template_path"] == "brand.potx"
    assert len(Presentation(tmp_path / "custom.pptx").slides) == 2


def test_animated_gradient_template_writes_native_transition_and_gradient(tmp_path):
    tool = make_build_presentation_tool(workspace=tmp_path)
    result = tool(
        title="Motion and gradient",
        slides=[{"title": "A chart-ready story", "bullets": ["12%", "24%", "36%"], "layout": "metric-grid"}],
        pptx_path="motion.pptx",
        pdf_path="motion.pdf",
        template_id="neon-flow",
    )
    assert result["ok"] is True
    assert result["template_id"] == "neon-flow"
    with zipfile.ZipFile(tmp_path / "motion.pptx") as archive:
        slide_xml = archive.read("ppt/slides/slide2.xml")
        assert b"<p:transition" in slide_xml
        assert b"<p:push" in slide_xml
        assert b"<a:gradFill" in slide_xml


def test_reference_inspired_templates_use_distinct_cover_compositions(tmp_path):
    cover_image = tmp_path / "cover.png"
    _sample_image(cover_image)
    tool = make_build_presentation_tool(workspace=tmp_path)
    for template_id in ("science-studio", "environmental-fieldwork", "cyan-infographic"):
        result = tool(
            title=f"{template_id} visual system",
            subtitle="A distinct reference-inspired composition",
            slides=[{"title": "The content remains editable", "layout": "statement"}],
            pptx_path=f"{template_id}.pptx",
            pdf_path=f"{template_id}.pdf",
            template_id=template_id,
            cover_image_path="cover.png",
        )
        assert result["ok"] is True
        assert result["template_id"] == template_id
        assert result["images_embedded"] == 1
        presentation = Presentation(tmp_path / f"{template_id}.pptx")
        assert len(presentation.slides[0].shapes) >= 5
        assert len(presentation.slides[0].shapes._spTree.xpath(".//p:pic")) == 1


def test_presentation_typography_uses_readable_hierarchy(tmp_path):
    tool = make_build_presentation_tool(workspace=tmp_path)
    result = tool(
        title="A concise presentation title",
        subtitle="One supporting line",
        slides=[{"title": "A takeaway title", "takeaway": "A readable key message", "layout": "statement"}],
        pptx_path="typography.pptx",
        pdf_path="typography.pdf",
        template_id="atlas",
    )
    assert result["ok"] is True
    presentation = Presentation(tmp_path / "typography.pptx")
    cover_sizes = [
        paragraph.font.size.pt
        for shape in presentation.slides[0].shapes
        if getattr(shape, "has_text_frame", False)
        for paragraph in shape.text_frame.paragraphs
        if paragraph.font.size
    ]
    content_sizes = [
        paragraph.font.size.pt
        for shape in presentation.slides[1].shapes
        if getattr(shape, "has_text_frame", False)
        for paragraph in shape.text_frame.paragraphs
        if paragraph.font.size
    ]
    assert max(cover_sizes) >= 50
    assert max(content_sizes) >= 35


def test_quality_gate_rejects_markdown_renamed_as_pdf(tmp_path):
    reports = tmp_path / "reports"
    reports.mkdir()
    (reports / "deck.pptx").write_bytes(b"PK\x03\x04")
    (reports / "deck.pdf").write_text("# This is Markdown", encoding="utf-8")
    run = {
        "deliverable": "presentation",
        "artifact_paths": ["reports/deck.pptx", "reports/deck.pdf"],
        "method": "standard",
        "image_mode": "none",
    }
    quality = SessionManager._research_quality(run, workspace=tmp_path)
    assert quality["status"] == "needs_attention"
    assert "The PDF file is not a real PDF document." in quality["issues"]


def test_presentation_studio_skill_has_no_placeholders():
    loader = SkillLoader(
        [
            (
                __import__("pathlib").Path(__file__).resolve().parents[1]
                / "skills"
            )
        ]
    )
    skill = loader.get("presentation-studio")
    assert skill is not None
    assert "build_presentation" in skill.instructions
    assert "art-direction.md" in skill.instructions
    assert "TODO" not in skill.instructions
    art_direction = (
        __import__("pathlib").Path(__file__).resolve().parents[1]
        / "skills"
        / "presentation-studio"
        / "references"
        / "art-direction.md"
    ).read_text(encoding="utf-8")
    assert "Never repeat the same layout more than twice consecutively." in art_direction
    assert "Deck title: 50–64 pt" in art_direction
